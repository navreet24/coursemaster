import requests
from flask import Flask, request, jsonify
from io import BytesIO
import PyPDF2
from docx import Document
from datetime import datetime
import time
from selenium import webdriver
from selenium.webdriver.chrome.service import Service
from selenium.webdriver.common.by import By
from selenium.webdriver.chrome.options import Options
from selenium.webdriver.support.ui import WebDriverWait
from selenium.webdriver.support import expected_conditions as EC
from webdriver_manager.chrome import ChromeDriverManager
from bs4 import BeautifulSoup
from pathlib import Path
from flask_cors import CORS
import os


# Initialize the Flask app
app = Flask(__name__)
CORS(app, resources={r"/*": {"origins": "http://localhost:5173"}})

# API Configuration
openrouter_api_key = "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"  # Replace with your API key
client_id = os.getenv("CLIENT_ID")
client_secret = os.getenv("CLIENT_SECRET")
tenant_id = os.getenv("TENANT_ID")

# Helper Functions
def get_access_token():
    token_url = f"https://login.microsoftonline.com/{tenant_id}/oauth2/v2.0/token"
    headers = {"Content-Type": "application/x-www-form-urlencoded"}
    body = {
        "grant_type": "client_credentials",
        "client_id": client_id,
        "client_secret": client_secret,
        "scope": "https://graph.microsoft.com/.default"
    }
    response = requests.post(token_url, headers=headers, data=body)
    if response.status_code == 200:
        return response.json().get('access_token')
    else:
        raise Exception(f"Failed to obtain access token: {response.text}")

def get_site_id(site_url, access_token):
    full_url = f'https://graph.microsoft.com/v1.0/sites/{site_url}'
    headers = {'Authorization': f'Bearer {access_token}'}
    response = requests.get(full_url, headers=headers)
    response.raise_for_status()
    return response.json().get('id')

def get_drive_id(site_id, access_token):
    drives_url = f'https://graph.microsoft.com/v1.0/sites/{site_id}/drives'
    headers = {'Authorization': f'Bearer {access_token}'}
    response = requests.get(drives_url, headers=headers)
    response.raise_for_status()
    drives = response.json().get('value', [])
    if not drives:
        raise ValueError("No drives found for the given site.")
    return drives[0]['id']

def list_folder_contents(site_id, drive_id, access_token, folder_id='root'):
    folder_contents_url = f'https://graph.microsoft.com/v1.0/sites/{site_id}/drives/{drive_id}/items/{folder_id}/children'
    headers = {'Authorization': f'Bearer {access_token}'}
    response = requests.get(folder_contents_url, headers=headers)
    response.raise_for_status()
    folder_contents = response.json()
    items_list = []
    if 'value' in folder_contents:
        for item in folder_contents['value']:
            if 'folder' in item:
                items_list.append((item['id'], item['name'], 'Folder'))
            elif 'file' in item:
                items_list.append((item['id'], item['name'], 'File'))
    return items_list

def download_and_extract_pdf_text(site_id, drive_id, file_id, access_token):
    download_url = f"https://graph.microsoft.com/v1.0/sites/{site_id}/drives/{drive_id}/items/{file_id}/content"
    headers = {'Authorization': f'Bearer {access_token}'}
    try:
        response = requests.get(download_url, headers=headers)
        response.raise_for_status()
        pdf_file = BytesIO(response.content)
        pdf_reader = PyPDF2.PdfReader(pdf_file)
        extracted_text = ""
        for page in pdf_reader.pages:
            extracted_text += page.extract_text()
        return extracted_text
    except Exception as e:
        print(f"Error processing PDF: {e}")
        return None

def chunk_text(text, chunk_size=3000):
    return [text[i:i + chunk_size] for i in range(0, len(text), chunk_size)]

def create_word_document(content, unit_name, filename_prefix="generated_document"):
    # Generate a unique filename with a timestamp
    #timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    #filename = f"{unit_name}_Formative_{timestamp}.docx"
    base_filename = f"{unit_name}_Formative.docx"
    file_path = Path(base_filename)
    

    # Handle versioning if the file already exists
    version = 1
    while file_path.exists():
        base_filename = f"{unit_name}_Formative_v{version}.docx"
        file_path = Path(base_filename)
        version +=1

    # Create the Word document
    doc = Document()
    doc.add_heading("Generated Assessment", level=1)
    for section in content.split("\n\n"):
        doc.add_paragraph(section)
    doc.save(base_filename)
    return base_filename

def find_folder_id_from_path(site_id, drive_id, folder_path, access_token):
    parts = folder_path.strip('/').split('/')
    folder_id = 'root'
    for part in parts:
        contents = list_folder_contents(site_id, drive_id, access_token, folder_id)
        next_folder_id = None
        for item_id, item_name, item_type in contents:
            if item_type == 'Folder' and item_name == part:
                next_folder_id = item_id
                break
        if next_folder_id is None:
            raise ValueError(f"Folder '{part}' not found in the path.")
        folder_id = next_folder_id
    return folder_id

def extract_data_under_heading(soup, heading_text):
    headings = soup.find_all(lambda tag: tag.name in ["h1", "h2", "h3", "h4"] and heading_text.lower() in tag.get_text(strip=True).lower())
    if headings:
        heading = headings[0]
        container = heading.find_next('div', {'class': 'table-std'}) or heading.find_next('div')
        if container:
            rows = container.find_all('tr')
            if rows:
                data = []
                for row in rows:
                    cols = row.find_all('td')
                    if cols:
                        data.append([col.get_text(strip=True) for col in cols])
                return data
            else:
                return container.get_text(strip=True)
        else:
            return None
    return None

def scrape_training_details(unit_code, retries=3):
    options = Options()
    options.add_argument("--headless")
    options.add_argument("--disable-gpu")
    options.add_argument("--no-sandbox")
    options.add_argument("--disable-dev-shm-usage")
    options.add_argument("--disable-extensions")
    options.add_argument("--disable-popup-blocking")
    options.add_argument("--incognito")

    attempt = 0
    while attempt < retries:
        driver = None
        try:
            # Initialize Selenium WebDriver
            driver = webdriver.Chrome(service=Service(ChromeDriverManager().install()), options=options)
            url = f"https://training.gov.au/training/details/{unit_code}/unitdetails"
            print(f"Attempt {attempt + 1}: Fetching {url}")
            driver.get(url)

            # Wait for the main content to load
            WebDriverWait(driver, 60).until(
                EC.presence_of_element_located((By.CSS_SELECTOR, 'div.table-std'))
            )

            # Parse the HTML content with BeautifulSoup
            html = driver.page_source
            soup = BeautifulSoup(html, 'html.parser')

            # Extract Elements and Performance Criteria
            elements_and_performance_criteria = extract_data_under_heading(soup, "Elements and performance criteria")

            # Extract Knowledge Evidence
            knowledge_evidence = extract_data_under_heading(soup, "Knowledge evidence")

            # Return the extracted details
            return knowledge_evidence, elements_and_performance_criteria

        except Exception as e:
            print(f"Error during attempt {attempt + 1}: {e}")
            # Save the page source for debugging
            if driver:
                html = driver.page_source
                debug_file = f"scraping_error_attempt_{attempt + 1}.html"
                with open(debug_file, "w", encoding="utf-8") as f:
                    f.write(html)
                print(f"Debug page source saved to {debug_file}")
            
            # Increment the attempt counter
            attempt += 1
            if attempt < retries:
                print(f"Retrying... (Attempt {attempt + 1} of {retries})")
            else:
                print("Max retries reached. Scraping failed.")

        finally:
            # Ensure WebDriver is closed
            if driver:
                driver.quit()

    # If all attempts fail, return None
    return None, None
def chunk_text_by_chapters(text):
    # Simple heuristic: Split by headings or some delimiter
    return text.split("Chapter ")  # Assuming "Chapter" appears in the learner guide



def upload_file_to_sharepoint(site_id, drive_id, folder_id, file_path, access_token):
    upload_url = f"https://graph.microsoft.com/v1.0/sites/{site_id}/drives/{drive_id}/items/{folder_id}:/{file_path.split('/')[-1]}:/content"
    headers = {
        'Authorization': f'Bearer {access_token}',
        'Content-Type': 'application/octet-stream'
    }
    with open(file_path, 'rb') as file:
        response = requests.put(upload_url, headers=headers, data=file)
    response.raise_for_status()
    return response.json()

@app.route('/generate_assessment', methods=['POST'])
def generate_assessment():
    try:
        data = request.get_json()
        print(f"Incoming request data: {data}")

        # Retrieve username and selected units
        Username = data.get('username')
        if not Username:
            return jsonify({'error': 'Username is required in the request payload'}), 400

        # Retrieve selected units from the frontend
        selected_units = data.get('selected_units', [])
        if not selected_units:
            return jsonify({'error': 'Selected units are required'}), 400

        # Use the first unit code or process as needed
        unit_code = selected_units[0] if selected_units else None
        if not unit_code:
            return jsonify({'error': 'Unit code is required'}), 400

        # Get the unit name and expected learner guide name
        unit_name = unit_code  # You can modify this to get the unit name if available
        expected_learner_guide_name = f"{unit_name}-LG-v1.0.pdf"

        site_url = "aoscaustralia.sharepoint.com:/sites/coursemaster"
        base_directory = "User directory"

        access_token = get_access_token()
        site_id = get_site_id(site_url, access_token)
        drive_id = get_drive_id(site_id, access_token)

        # Get the user's folder path
        user_directory_path = f"{base_directory}/{Username}"
        user_folder_id = find_folder_id_from_path(site_id, drive_id, user_directory_path, access_token)
        user_folder_contents = list_folder_contents(site_id, drive_id, access_token, user_folder_id)

        # Check if the expected learner guide is available in the user's folder
        learner_guide_file_id = None
        for item_id, item_name, item_type in user_folder_contents:
            if item_type == 'File' and item_name == expected_learner_guide_name:
                learner_guide_file_id = item_id
                break

        if not learner_guide_file_id:
            return jsonify({'error': f"Learner guide '{expected_learner_guide_name}' is not available in your folder."}), 404

        # If learner guide is found, proceed with extraction
        pdf_content = download_and_extract_pdf_text(site_id, drive_id, learner_guide_file_id, access_token)
        if not pdf_content:
            return jsonify({'error': 'Failed to extract text from the learner guide PDF.'}), 500

        # Scrape additional details
        knowledge_evidence, performance_criteria = scrape_training_details(unit_code)

        # Updated Prompt for Each Chapter
        chapters = chunk_text_by_chapters(pdf_content)  # Assuming a function to split content into chapters
        assessments = []

        # Define question types for chapters
        question_types = [
            {"type": "MCQs", "count": 5},
            {"type": "True/False", "count": 5},
            {"type": "Short Answer", "count": 5},
            {"type": "Long Answer", "count": 5},
            {"type": "Fill in the Blanks", "count": 5},
            {"type": "Match the Following", "count": 5},
        ]

        for chapter_idx, chapter_content in enumerate(chapters):
            questions = []
            chapter_title = f"Chapter {chapter_idx + 1}"
            question_type = question_types[chapter_idx % len(question_types)]  # Rotate question types for variety

            # Prepare content for the API
            messages = [
                {"role": "system", "content": "You are an educational assistant that generates assessments with specific question types based on provided learner guides, knowledge evidence, and performance criteria."},
                {"role": "user", "content": f"Please generate {question_type['count']} questions of type '{question_type['type']}' for {chapter_title}. "
                                             f"Base the questions on the following:\n\n"
                                             f"Learner Guide:\n{chapter_content}\n\n"
                                             f"Knowledge Evidence:\n{knowledge_evidence}\n\n"
                                             f"Performance Criteria:\n{performance_criteria}\n\n"
                                             f"Include detailed answers for each question."}
            ]

            while True:
                try:
                    response = requests.post(
                        "https://openrouter.ai/api/v1/chat/completions",
                        headers={
                            "Authorization": f"Bearer {openrouter_api_key}",
                            "Content-Type": "application/json",
                        },
                        json={
                            "model": "meta-llama/llama-3.2-3b-instruct:free",
                            "messages": messages,
                            "max_tokens": 10000,
                            "temperature": 0.7,
                        },
                    )

                    if response.status_code == 429:
                        retry_after = int(response.headers.get("X-RateLimit-Reset", time.time())) - int(time.time())
                        print(f"Rate limit exceeded. Retrying after {retry_after} seconds...")
                        time.sleep(retry_after)
                    else:
                        break
                except requests.exceptions.RequestException as e:
                    print(f"Error making request: {e}")
                    return jsonify({'error': f"Request error: {e}"}), 500

            try:
                response_data = response.json()

                if 'choices' in response_data:
                    chapter_questions = response_data['choices'][0]['message']['content'].strip()
                    questions.append(chapter_questions)
                else:
                    questions.append(f"Failed to generate questions for {chapter_title}. No 'choices' in the response.")
            except (KeyError, IndexError) as e:
                print(f"Error parsing API response: {e}")
                questions.append(f"Failed to parse API response for {chapter_title}.")

            assessments.append(f"{chapter_title}\n\n{'\n'.join(questions)}")

        combined_assessment = "\n\n".join(assessments)
        word_filename = create_word_document(combined_assessment, unit_name)

        # Upload the Word document to the user's folder
        try:
            upload_file_to_sharepoint(site_id, drive_id, user_folder_id, word_filename, access_token)
            return jsonify({'assessment': combined_assessment, 'word_file': word_filename, 'upload_status': 'Success'}), 200
        except Exception as upload_error:
            return jsonify({'assessment': combined_assessment, 'word_file': word_filename, 'upload_status': 'Failed', 'error': str(upload_error)}), 500
    except Exception as e:
        return jsonify({'error': str(e)}), 500


if __name__ == '__main__':
    app.run(debug=True, port=5012)
