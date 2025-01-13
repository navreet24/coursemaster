from selenium import webdriver
from selenium.webdriver.chrome.service import Service
from selenium.webdriver.common.by import By
from selenium.webdriver.chrome.options import Options
from selenium.webdriver.support.ui import WebDriverWait
from selenium.webdriver.support import expected_conditions as EC
from webdriver_manager.chrome import ChromeDriverManager
from bs4 import BeautifulSoup
from docx import Document

# Function to construct the URL dynamically
def construct_url(unit_code):
    base_url = "https://training.gov.au/training/details/"
    return f"{base_url}{unit_code}"

# Function to extract data under a specific heading
def extract_data_under_heading(soup, heading_text):
    # Look for the heading with a broader search criteria
    headings = soup.find_all(lambda tag: tag.name in ["h1", "h2", "h3", "h4"] and heading_text.lower() in tag.get_text(strip=True).lower())
    if headings:
        for heading in headings:
            print(f"Found heading: {heading.get_text(strip=True)}")  # Debugging print for headings
        # Extract the content under the first matching heading
        heading = headings[0]
        print(f"Heading '{heading_text}' found!")
        container = heading.find_next('div', {'class': 'table-std'}) or heading.find_next('div')
        if container:
            rows = container.find_all('tr')
            if rows:
                data = []
                for row in rows:
                    cols = row.find_all('td')
                    if cols:
                        data.append([col.get_text(strip=True) for col in cols])
                return data
            else:
                return container.get_text(strip=True)
        else:
            print(f"No content found under '{heading_text}'.")
    else:
        print(f"Heading '{heading_text}' not found!")
    return None

# Selenium setup
options = Options()
options.headless = True
driver = webdriver.Chrome(service=Service(ChromeDriverManager().install()), options=options)

# Get the unit code dynamically (user input or pass the unit code directly)
unit_code = input("Enter the unit code: ").strip()

# Construct the URL based on the unit code
url = construct_url(unit_code)
driver.get(url)

# Increase the wait time if necessary to ensure the content is fully loaded
try:
    WebDriverWait(driver, 30).until(
        EC.presence_of_element_located((By.CSS_SELECTOR, 'div.table-std'))
    )
except Exception as e:
    print("Timeout error:", e)

# Get page source
html = driver.page_source

# Debugging: Print the entire HTML structure to inspect the page
print(html)  # You can remove this line if the HTML is too large

soup = BeautifulSoup(html, 'html.parser')

# Debugging: Print first 10 divs to inspect the structure
divs = soup.find_all('div')
for div in divs[:10]:  # Print the first 10 divs
    print(div.prettify())  # Pretty print each div

# Extract data under 'Elements and Performance Criteria'
elements_and_performance_criteria = extract_data_under_heading(soup, "Elements and Performance Criteria")

# Extract data under 'Knowledge Evidence'
knowledge_evidence = extract_data_under_heading(soup, "Knowledge Evidence")

# Close the driver
driver.quit()

# Create a Word document
document = Document()

# Add a title to the document
document.add_heading("Training Unit Details", level=1)

# Add Elements and Performance Criteria to the document
if elements_and_performance_criteria:
    document.add_heading("Elements and Performance Criteria", level=2)
    for row in elements_and_performance_criteria:
        document.add_paragraph(", ".join(row))  # Combine row data into a single line

# Add Knowledge Evidence to the document
if knowledge_evidence:
    document.add_heading("Knowledge Evidence", level=2)
    if isinstance(knowledge_evidence, list):
        for row in knowledge_evidence:
            document.add_paragraph(", ".join(row))  # Combine row data into a single line
    else:
        document.add_paragraph(knowledge_evidence)

# Save the document
output_file = f"{unit_code}_Training_Unit_Details.docx"
document.save(output_file)
print(f"Document saved as {output_file}")
