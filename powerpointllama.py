import requests
from flask import Flask, request, jsonify
from io import BytesIO
from pptx import Presentation
from pptx.util import Pt
import os
from flask_cors import CORS
from docx import Document

app = Flask(__name__)
CORS(app, resources={r"/*": {"origins": "http://localhost:5174"}})

# Configuration
META_LLAMA_API_KEY = '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'
CLIENT_ID = 'b25fac3a-8c66-46af-b161-840828c538f3'
CLIENT_SECRET = 'GFm8Q~q6lG0UO3_ZT-~LVMNjgvEND6~cRXt0pc_f'
TENANT_ID = '4dab0fef-f02d-440b-97c3-712e9483bd68'

def get_access_token():
    url = f"https://login.microsoftonline.com/{TENANT_ID}/oauth2/v2.0/token"
    headers = {"Content-Type": "application/x-www-form-urlencoded"}
    data = {
        "grant_type": "client_credentials",
        "client_id": CLIENT_ID,
        "client_secret": CLIENT_SECRET,
        "scope": "https://graph.microsoft.com/.default"
    }
    response = requests.post(url, headers=headers, data=data)
    response.raise_for_status()
    return response.json()['access_token']

def get_site_id(site_url, access_token):
    url = f'https://graph.microsoft.com/v1.0/sites/{site_url}'
    headers = {'Authorization': f'Bearer {access_token}'}
    response = requests.get(url, headers=headers)
    response.raise_for_status()
    return response.json()['id']

def get_drive_id(site_id, access_token):
    url = f'https://graph.microsoft.com/v1.0/sites/{site_id}/drives'
    headers = {'Authorization': f'Bearer {access_token}'}
    response = requests.get(url, headers=headers)
    response.raise_for_status()
    return response.json()['value'][0]['id']

def find_folder_id(site_id, drive_id, folder_path, access_token):
    folder_id = 'root'
    for folder in folder_path.strip('/').split('/'):
        url = f'https://graph.microsoft.com/v1.0/sites/{site_id}/drives/{drive_id}/items/{folder_id}/children'
        headers = {'Authorization': f'Bearer {access_token}'}
        response = requests.get(url, headers=headers)
        response.raise_for_status()
        for item in response.json().get('value', []):
            if item.get('folder') and item['name'] == folder:
                folder_id = item['id']
                break
        else:
            raise ValueError(f"Folder '{folder}' not found in path.")
    return folder_id

def extract_word_text(file_content):
    document = Document(BytesIO(file_content))
    return "\n".join(paragraph.text for paragraph in document.paragraphs)

def generate_presentation(chunks, title):
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[0])
    slide.shapes.title.text = title

    slide_limit = 50
    for i, chunk in enumerate(chunks[:slide_limit]):
        slide = prs.slides.add_slide(prs.slide_layouts[1])
        slide.shapes.title.text = f"Section {i + 1}"

        content = slide.placeholders[1]
        for paragraph_text in chunk.split("\n"):
            p = content.text_frame.add_paragraph()
            p.text = paragraph_text.strip()
            p.font.size = Pt(14)
    return prs

@app.route('/generate_presentation', methods=['POST'])
def generate_presentation_route():
    try:
        data = request.json
        site_url = "aoscaustralia.sharepoint.com:/sites/coursemaster"
        read_folder = "learner guides"
        save_folder = f"User directory/{data.get('username')}"
        access_token = get_access_token()

        site_id = get_site_id(site_url, access_token)
        drive_id = get_drive_id(site_id, access_token)
        read_folder_id = find_folder_id(site_id, drive_id, read_folder, access_token)
        save_folder_id = find_folder_id(site_id, drive_id, save_folder, access_token)

        url = f'https://graph.microsoft.com/v1.0/sites/{site_id}/drives/{drive_id}/items/{read_folder_id}/children'
        headers = {'Authorization': f'Bearer {access_token}'}
        response = requests.get(url, headers=headers)
        response.raise_for_status()

        for item in response.json().get('value', []):
            if item['name'].endswith('.docx'):
                file_id = item['id']
                download_url = f"https://graph.microsoft.com/v1.0/sites/{site_id}/drives/{drive_id}/items/{file_id}/content"
                doc_response = requests.get(download_url, headers=headers)
                doc_response.raise_for_status()

                text = extract_word_text(doc_response.content)
                chunks = [text[i:i + 1000] for i in range(0, len(text), 1000)]
                presentation = generate_presentation(chunks, item['name'].rsplit('.', 1)[0])
                
                output_file = f"{item['name'].rsplit('.', 1)[0]}.pptx"
                presentation.save(output_file)

                upload_url = f"https://graph.microsoft.com/v1.0/sites/{site_id}/drives/{drive_id}/items/{save_folder_id}:/{output_file}:/content"
                with open(output_file, 'rb') as file_data:
                    upload_response = requests.put(upload_url, headers=headers, data=file_data)
                    upload_response.raise_for_status()
                os.remove(output_file)

                return jsonify({"message": "Presentation generated and uploaded."}), 200

        return jsonify({"error": "No Word documents found."}), 404

    except Exception as e:
        return jsonify({"error": str(e)}), 500

if __name__ == '__main__':
    app.run(debug=True, port=5008)
